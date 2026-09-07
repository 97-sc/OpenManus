"""生成一份示例 Excel 财报（用于测试 FinancialReportTool）"""
import os
from openpyxl import Workbook

OUT_DIR = r"D:\projects\OpenManus\examples"
OUT_FILE = os.path.join(OUT_DIR, "tesco_fy2024_sample.xlsx")

os.makedirs(OUT_DIR, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "财报数据"

# 第一行：期数（年份）
ws.cell(row=1, column=1, value="科目")
ws.cell(row=1, column=2, value="2024")
ws.cell(row=1, column=3, value="2023")
ws.cell(row=1, column=4, value="2022")

# 后续行：科目 + 各期数据
data = [
    ("营业收入", 1000000, 900000, 800000),
    ("营业成本", 700000, 650000, 600000),
    ("净利润", 100000, 80000, 60000),
    ("总资产", 1500000, 1300000, 1200000),
    ("总负债", 900000, 800000, 750000),
    ("股东权益", 600000, 500000, 450000),
    ("流动资产", 600000, 550000, 500000),
    ("流动负债", 400000, 380000, 350000),
    ("存货", 150000, 140000, 130000),
]

for idx, row_data in enumerate(data, start=2):
    for col_idx, val in enumerate(row_data, start=1):
        ws.cell(row=idx, column=col_idx, value=val)

wb.save(OUT_FILE)
print(f"✅ Excel 已生成: {OUT_FILE}")
